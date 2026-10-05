"""Genera la carga inicial de casos a partir de los planes de trabajo en Excel.

Salidas:
  supabase/seed.sql     inserta fiscales y casos (no pisa casos que ya existan)
  src/data/demo.json    los mismos datos para el modo demo de la app

Uso:
  python3 -m pip install openpyxl
  python3 scripts/generar_seed.py [carpeta_salidas]

carpeta_salidas por defecto: ../salidas (junto a esta app en GestionAtencionCasos).
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
SALIDAS = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ.parent / "salidas"

FISCALES = [
    {
        "id": "11111111-1111-4111-8111-111111111111",
        "nombre": "García Vera, Noé Humberto",
        "archivo": "Plan_trabajo_5_meses.xlsx",
    },
    {
        "id": "22222222-2222-4222-8222-222222222222",
        "nombre": "Villanueva Flores, David Manuel",
        "archivo": "Plan_trabajo_5_meses_David.xlsx",
    },
]


def fecha(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return None


def leer(fiscal):
    ws = openpyxl.load_workbook(SALIDAS / fiscal["archivo"], data_only=True)["Plan por caso"]
    encabezado = [c.value for c in ws[5]]
    col = {nombre: i for i, nombre in enumerate(encabezado)}
    casos = []
    for fila in ws.iter_rows(min_row=6, values_only=True):
        if fila[col["Orden"]] is None:
            continue
        etapa = fila[col["Etapa actual"]]
        casos.append({
            "codigo": fila[col["Caso"]],
            "fiscal_id": fiscal["id"],
            "n_reporte": str(fila[col["N° reporte"]]),
            "etapa": etapa,
            "estado": fila[col["Estado"]],
            "fecha_inicio_etapa": fecha(fila[col["Inicio etapa"]]),
            "plazo_dias": fila[col["Plazo (días)"]],
            "tipo_dias": "habiles" if etapa == "CALIFICACION" else "calendario",
            "horas_estimadas": fila[col["Horas estimadas"]],
            "plan_orden": fila[col["Orden"]],
            "plan_inicio": fecha(fila[col["Inicio previsto"]]),
            "plan_fin": fecha(fila[col["Fin previsto"]]),
            "apoyo_otro_equipo": fila[col["Con apoyo del otro equipo"]] == "Sí",
        })
    return casos


def sql_valor(v):
    if v is None:
        return "null"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, (int, float)):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def main():
    todos = []
    for f in FISCALES:
        casos = leer(f)
        print(f"{f['nombre']}: {len(casos)} casos")
        todos.extend(casos)

    columnas = list(todos[0].keys())
    lineas = [
        "-- Carga inicial generada por scripts/generar_seed.py",
        "-- Fuente: planes de trabajo de 5 meses (reportes de plazos al 25/09/2026)",
        "",
        "insert into public.fiscales (id, nombre) values",
        ",\n".join(f"  ({sql_valor(f['id'])}, {sql_valor(f['nombre'])})" for f in FISCALES),
        "on conflict (id) do nothing;",
        "",
        f"insert into public.casos ({', '.join(columnas)}) values",
        ",\n".join("  (" + ", ".join(sql_valor(c[k]) for k in columnas) + ")" for c in todos),
        "on conflict (codigo) do nothing;",
        "",
    ]
    (RAIZ / "supabase" / "seed.sql").write_text("\n".join(lineas), encoding="utf-8")

    demo = {
        "generado": date.today().isoformat(),
        "fiscales": [{"id": f["id"], "nombre": f["nombre"], "horas_equipo_dia": 24} for f in FISCALES],
        "casos": todos,
    }
    (RAIZ / "src" / "data" / "demo.json").write_text(json.dumps(demo, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"Total: {len(todos)} casos -> supabase/seed.sql y src/data/demo.json")


if __name__ == "__main__":
    main()
